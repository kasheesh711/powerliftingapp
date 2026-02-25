#!/usr/bin/env python3
"""Export locked golden fixture files from workbook DB snapshot sheets."""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / 'Kev Ultimate Comeback.xlsx'
OUT_DIR = ROOT / 'docs' / 'fixtures' / 'golden'
TARGET_BLOCKS = ['Block 2', 'Block 3', 'Block 4 (2026)', 'Block 5', 'Block 3 (Data)']


def load_db_rows(workbook: openpyxl.Workbook) -> dict[str, list[dict[str, object]]]:
    rows = {name: [] for name in TARGET_BLOCKS}
    if '_DB_Blocks' not in workbook.sheetnames:
        return rows

    ws = workbook['_DB_Blocks']
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]

    for r in range(2, ws.max_row + 1):
        values = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        block_name = str(values[0]) if values and values[0] is not None else ''
        if block_name not in rows:
            continue

        row = {
            str(headers[i] if headers[i] is not None else f'col_{i+1}'): values[i]
            for i in range(len(headers))
        }
        rows[block_name].append(row)

    return rows


def load_db_metadata(workbook: openpyxl.Workbook) -> dict[str, dict[str, object]]:
    out = {name: {} for name in TARGET_BLOCKS}
    if '_DB_Metadata' not in workbook.sheetnames:
        return out

    ws = workbook['_DB_Metadata']
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]

    for r in range(2, ws.max_row + 1):
        values = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        block_name = str(values[0]) if values and values[0] is not None else ''
        if block_name not in out:
            continue

        record = {
            str(headers[i] if headers[i] is not None else f'col_{i+1}'): values[i]
            for i in range(len(headers))
        }

        for key in ['RecapsJSON', 'PeakE1RMsJSON', 'ParserReportJSON']:
            raw = record.get(key)
            if not isinstance(raw, str):
                continue
            try:
                record[key.replace('JSON', '')] = json.loads(raw)
            except json.JSONDecodeError:
                record[key.replace('JSON', '')] = raw

        out[block_name] = record

    return out


def sanitize_filename(name: str) -> str:
    return (
        name.replace(' ', '_')
        .replace('(', '')
        .replace(')', '')
        .replace('-', '_')
        .replace('__', '_')
    )


def main() -> None:
    workbook = openpyxl.load_workbook(WORKBOOK, data_only=True)
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    db_rows = load_db_rows(workbook)
    db_meta = load_db_metadata(workbook)

    for block in TARGET_BLOCKS:
        payload = {
            'blockName': block,
            'sheetExists': block in workbook.sheetnames,
            'sheetDimensions': {
                'rows': workbook[block].max_row if block in workbook.sheetnames else 0,
                'cols': workbook[block].max_column if block in workbook.sheetnames else 0,
            },
            'dbRowCount': len(db_rows.get(block, [])),
            'dbRows': db_rows.get(block, []),
            'metadata': db_meta.get(block, {}),
        }

        target = OUT_DIR / f"{sanitize_filename(block)}.json"
        with target.open('w', encoding='utf-8') as handle:
            json.dump(payload, handle, indent=2, ensure_ascii=True)

    print(f'Golden fixtures exported to {OUT_DIR}')


if __name__ == '__main__':
    main()
